"""http://127.0.0.1:5050"""

import os
import sqlite3
from datetime import date, datetime, timedelta

from flask import Flask, g, jsonify, request, render_template, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import io

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.environ.get("TRACKER_DB_PATH") or os.path.join(BASE_DIR, "database", "tracker.db")
SCHEMA_PATH = os.path.join(BASE_DIR, "database", "schema.sql")

app = Flask(__name__)

# List of animal profile cards that I have personally chosen. Feel free to change to your liking
ANIMAL_NAMES = [
    "German Shepherd", "Doberman Pinscher", "Husky", "Pomeranian", "Samoyed",
    "Toucan", "Parrot", "Dove", "Pigeon", "Owl", "Dolphin", "Shark", "Orca", "Narwhal", "Octopus",
    "Jellyfish", "Lion", "Tiger", "Bear", "Wolf", "Artic Fox", "Grey Wolf", "Deer", "Rabbit", 
    "Dragon", "Unicorn", "Pegasus", "Phoenix", "Griffin", "Mouse", "Snake", "Kangaroo", "Raccoon",
    "Cheetah", "Zebra", "Giraffe", "Panther", "Jaguar", "Sheep", "Goat", "Alpaca", "Capybara",
]

def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
    return g.db


@app.teardown_appcontext
def close_db(exception=None):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def create_db():
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    with open(SCHEMA_PATH, "r") as f:
        conn.executescript(f.read())
    conn.commit()
    conn.close()


def today_str():
    return date.today().isoformat()


def ensure_daily_records(db, class_id, record_date):
    students = db.execute(
        "SELECT id FROM students WHERE class_id = ?", (class_id,)
    ).fetchall()
    for s in students:
        db.execute(
            """INSERT OR IGNORE INTO daily_records (student_id, record_date)
               VALUES (?, ?)""",
            (s["id"], record_date),
        )
    db.commit()


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/classes", methods=["GET"])
def list_classes():
    db = get_db()
    rows = db.execute(
        "SELECT id, name, student_count FROM classes ORDER BY created_at DESC"
    ).fetchall()
    result = []
    for r in rows:
        result.append(dict(r))
    return jsonify(result)

@app.route("/api/classes", methods=["POST"])
def create_class():
    data = request.get_json(force=True)
    name = (data.get("name") or "").strip() or "Untitled Class"
    try:
        count = int(data.get("student_count", 0))
    except (TypeError, ValueError):
        count = 0

    if count < 1 or count > len(ANIMAL_NAMES):
        return jsonify({
            "error": f"Student count must be between 1 and {len(ANIMAL_NAMES)}."
        }), 400

    db = get_db()
    cur = db.execute(
        "INSERT INTO classes (name, student_count) VALUES (?, ?)",
        (name, count),
    )
    class_id = cur.lastrowid

    for i in range(count):
        db.execute(
            "INSERT INTO students (class_id, animal_name, seat_order) VALUES (?, ?, ?)",
            (class_id, ANIMAL_NAMES[i], i + 1),
        )
    db.commit()

    ensure_daily_records(db, class_id, today_str())

    return jsonify({"id": class_id, "name": name, "student_count": count}), 201


@app.route("/api/dashboard/<int:class_id>", methods=["GET"])
def get_dashboard(class_id):
    db = get_db()
    today = today_str()
    ensure_daily_records(db, class_id, today)

    rows = db.execute(
        """SELECT s.id AS student_id, s.animal_name, s.seat_order,
                  d.assigned_homework, d.completed_homework, d.meets_expectations
           FROM students s
           JOIN daily_records d ON d.student_id = s.id AND d.record_date = ?
           WHERE s.class_id = ?
           ORDER BY s.seat_order""",
        (today, class_id),
    ).fetchall()

    students_list = []
    for r in rows:
        students_list.append(dict(r))
    return jsonify({"date": today, "students": students_list})


VALID_FIELDS = {"assigned_homework", "completed_homework", "meets_expectations"}


@app.route("/api/update", methods=["POST"])
def update_record():
    data = request.get_json(force=True)
    student_id = data.get("student_id")
    field = data.get("field")
    value = 1 if data.get("value") else 0

    if field not in VALID_FIELDS or not isinstance(student_id, int):
        return jsonify({"error": "Invalid field or student_id."}), 400

    db = get_db()
    today = today_str()
    db.execute(
        f"""UPDATE daily_records SET {field} = ?
            WHERE student_id = ? AND record_date = ?""",
        (value, student_id, today),
    )
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/reports/<int:student_id>", methods=["GET"])
def get_reports(student_id):
    db = get_db()
    rows = db.execute(
        """SELECT id, report_text, created_at FROM reports
           WHERE student_id = ? ORDER BY created_at DESC""",
        (student_id,),
    ).fetchall()
    result = []
    for r in rows:
        result.append(dict(r))
    return jsonify(result)


@app.route("/api/reports", methods=["POST"])
def add_report():
    data = request.get_json(force=True)
    student_id = data.get("student_id")
    report_text = (data.get("report_text") or "").strip()

    if not isinstance(student_id, int) or not report_text:
        return jsonify({"error": "student_id and report_text are required."}), 400

    db = get_db()
    created_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    cur = db.execute(
        "INSERT INTO reports (student_id, report_text, created_at) VALUES (?, ?, ?)",
        (student_id, report_text, created_at),
    )
    db.commit()
    row = db.execute(
        "SELECT id, report_text, created_at FROM reports WHERE id = ?",
        (cur.lastrowid,),
    ).fetchone()
    return jsonify(dict(row)), 201


HEADER_FILL = PatternFill(start_color="3A5A40", end_color="3A5A40", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


@app.route("/api/export/<int:class_id>", methods=["GET"])
def export_week(class_id):
    db = get_db()
    cls = db.execute("SELECT * FROM classes WHERE id = ?", (class_id,)).fetchone()
    if not cls:
        return jsonify({"error": "Class not found."}), 404

    end = date.today()
    start = end - timedelta(days=6)

    students = db.execute(
        "SELECT id, animal_name FROM students WHERE class_id = ? ORDER BY seat_order",
        (class_id,),
    ).fetchall()

    reports = db.execute(
        """SELECT n.student_id, n.report_text, n.created_at FROM reports n
           JOIN students s ON s.id = n.student_id
           WHERE s.class_id = ? AND date(n.created_at) BETWEEN ? AND ?
           ORDER BY n.created_at ASC""",
        (class_id, start.isoformat(), end.isoformat()),
    ).fetchall()

    wb = Workbook()

    ws = wb.active
    ws.title = "Behavior Reports"
    for col, text in enumerate(["Student", "Report", "Date"], start=1):
        cell = ws.cell(row=1, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 20
    ws.freeze_panes = "A2"

    id_to_name = {s["id"]: s["animal_name"] for s in students}
    for r_idx, n in enumerate(reports, start=2):
        ws.cell(row=r_idx, column=1, value=id_to_name.get(n["student_id"], "?")).font = Font(bold=True)
        ws.cell(row=r_idx, column=2, value=n["report_text"]).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r_idx, column=3, value=n["created_at"])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"{cls['name'].replace(' ', '_')}_{start.isoformat()}_to_{end.isoformat()}.xlsx"
    return send_file(
        buf,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


if __name__ == "__main__":
    create_db()
    app.run(debug=True, port=5050)
else:
    if not os.path.exists(DB_PATH):
        create_db()
